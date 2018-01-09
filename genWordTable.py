from docx import Document
from docx.shared import Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
import datetime

first_day = "2018-01-01"
d = datetime.datetime.strptime(first_day,"%Y-%m-%d")
delta = datetime.timedelta(days = 6 - d.weekday())

document = Document()
# styles = document.styles
# for s in styles:
# 	if s.type == WD_STYLE_TYPE.TABLE:
# 			document.add_paragraph("Table style is : " + s.name)
# 			document.add_table(3, 3, style = s)
# 			document.add_paragraph("\n")


table = document.add_table(rows = 53, cols = 4, style = "Light Grid Accent 1")
table.alignment = WD_TABLE_ALIGNMENT.CENTER

heading_cells = table.rows[0].cells
heading_cells[0].text = '2018'
heading_cells[1].text = '书'
heading_cells[2].text = '菜'
heading_cells[3].text = '电影'

first_col = table.columns[0].cells

for i in range(1, len(table.rows)):
	r = "(" + d.strftime("%m/%d") + '--' + (d + delta).strftime("%m/%d") + ")"
	first_col[i].text = "第%d周\n" % i + r
	d = d + datetime.timedelta(days = 7)
	delta = datetime.timedelta(days = 6)


document.add_page_break()
document.save('demo.docx')
